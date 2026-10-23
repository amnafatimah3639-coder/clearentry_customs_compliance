

import re
from pathlib import Path
from typing import List

from pydantic import BaseModel, Field
from crewai.tools import BaseTool
from pypdf import PdfReader
from docx import Document


class ReaderInput(BaseModel):
    file_path: str = Field(description="Path to the document to read.")


class CompareInput(BaseModel):
    invoice_text: str
    packing_list_text: str


class RegulationInput(BaseModel):
    query: str


class ReportInput(BaseModel):
    findings: str


def read_document(file_path: str) -> str:
    path = Path(file_path)
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        reader = PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    if suffix == ".docx":
        doc = Document(str(path))
        content = [p.text for p in doc.paragraphs]

        for table in doc.tables:
            for row in table.rows:
                content.append(" | ".join(cell.text for cell in row.cells))

        return "\n".join(content)

    if suffix == ".txt":
        return path.read_text(encoding="utf-8", errors="replace")

    raise ValueError(f"Unsupported file type: {suffix}")


class DocumentReaderTool(BaseTool):
    name: str = "Document Reader"
    description: str = (
        "Read a PDF, DOCX or TXT file. Supply its file_path."
    )
    args_schema: type[BaseModel] = ReaderInput

    def _run(self, file_path: str) -> str:
        text = read_document(file_path)

        if not text.strip():
            return "No selectable text found. The PDF may require OCR."

        return text[:24000]


class CrossDocumentCompareTool(BaseTool):
    name: str = "Cross Document Compare"
    description: str = (
        "Scan invoice and packing list for candidate dates, quantities, "
        "weights and currency amounts. Supply both document texts."
    )
    args_schema: type[BaseModel] = CompareInput

    def _run(self, invoice_text: str, packing_list_text: str) -> str:
        patterns = {
            "dates": r"\b\d{4}-\d{2}-\d{2}\b",
            "currencies": (
                r"\b(?:USD|EUR|GBP|PKR|CNY|AED|SAR)\s?"
                r"[\d,]+(?:\.\d{1,2})?\b"
            ),
            "weights": (
                r"\b\d+(?:,\d{3})*(?:\.\d+)?\s?"
                r"(?:kg|kgs|grams?|tonnes?|mt)\b"
            ),
            "quantities": (
                r"\b\d+(?:,\d{3})*(?:\.\d+)?\s?"
                r"(?:pcs?|pieces?|units?|cartons?|boxes?|sets?)\b"
            ),
        }

        results = []

        for label, pattern in patterns.items():
            invoice_values = re.findall(pattern, invoice_text, re.I)
            packing_values = re.findall(pattern, packing_list_text, re.I)

            results.append(
                f"{label.title()}\n"
                f"Invoice: {invoice_values[:30] or 'None found'}\n"
                f"Packing list: {packing_values[:30] or 'None found'}"
            )

        results.append(
            "These are candidate matches, not proof of a discrepancy. "
            "Compare labels, units and surrounding text."
        )

        return "\n\n".join(results)


class RegulationSearchTool(BaseTool):
    name: str = "Regulation Library Search"
    description: str = (
        "Search user-supplied regulation passages using a plain-language query."
    )
    args_schema: type[BaseModel] = RegulationInput

    source_names: List[str] = Field(default_factory=list)
    source_texts: List[str] = Field(default_factory=list)

    def _run(self, query: str) -> str:
        terms = set(re.findall(r"[a-zA-Z0-9]{3,}", query.lower()))
        matches = []

        for name, text in zip(self.source_names, self.source_texts):
            paragraphs = [
                p.strip()
                for p in re.split(r"\n+|(?<=[.!?])\s+", text)
                if p.strip()
            ]

            for paragraph in paragraphs:
                words = set(re.findall(r"[a-zA-Z0-9]{3,}", paragraph.lower()))
                score = len(terms & words)

                if score:
                    matches.append((score, name, paragraph))

        matches.sort(key=lambda item: item[0], reverse=True)

        if not matches:
            return (
                "No matching passage found in the supplied sources. "
                "This does not establish that no legal requirement exists."
            )

        return "\n\n".join(
            f"Source: {name}\nRelevance score: {score}\n{paragraph[:1200]}"
            for score, name, paragraph in matches[:8]
        )


class ReportStructureTool(BaseTool):
    name: str = "Compliance Report Structure"
    description: str = "Organize findings into a compliance report."
    args_schema: type[BaseModel] = ReportInput

    def _run(self, findings: str) -> str:
        return (
            "# Preliminary Customs Compliance Report\n\n"
            "## Executive summary\n"
            "## High-priority issues\n"
            "## Medium-priority issues\n"
            "## Low-priority verification notes\n"
            "## Cross-document evidence\n"
            "## Retrieved regulation passages and sources\n"
            "## Missing information and recommended follow-up\n"
            "## Limitations and human verification\n\n"
            "Draft findings:\n"
            f"{findings[:14000]}\n\n"
            "Do not certify legal compliance or invent citations."
        )
